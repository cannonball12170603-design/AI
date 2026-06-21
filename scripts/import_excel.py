"""
既存のExcel家計簿(.xlsx)をSupabaseに取り込むスクリプト。

事前準備:
  pip install openpyxl requests
  環境変数 SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY / SUPABASE_USER_ID を設定

使い方:
  python scripts/import_excel.py /path/to/家計簿.xlsx
"""
import os
import sys
import csv
import requests
import openpyxl

SUPABASE_URL = os.environ.get("SUPABASE_URL")
SERVICE_KEY = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
USER_ID = os.environ.get("SUPABASE_USER_ID")

# 生活管理シートの「収入合計」より上が収入、「支出合計」までが支出
INCOME_LABELS_STOP = "収入合計"
EXPENSE_LABELS_STOP = "支出合計"
SKIP_LABELS = {"繰越", "収入合計", "支出合計", "差し引き", "前月比", "楽天銀行", "株＋貯金"}


def insert_rows(table: str, rows: list[dict]):
    if not rows:
        return
    headers = {
        "apikey": SERVICE_KEY,
        "Authorization": f"Bearer {SERVICE_KEY}",
        "Content-Type": "application/json",
        "Prefer": "return=minimal",
    }
    resp = requests.post(f"{SUPABASE_URL}/rest/v1/{table}", json=rows, headers=headers)
    resp.raise_for_status()
    print(f"  -> {len(rows)} rows inserted into {table}")


def import_monthly_entries(ws):
    years = [c.value for c in ws[1]]
    months = [c.value for c in ws[2]]
    rows = []
    kind = "income"
    for r in range(3, 44):
        label = ws.cell(row=r, column=1).value
        if label is None:
            continue
        if label == INCOME_LABELS_STOP:
            kind = "expense"
            continue
        if label == EXPENSE_LABELS_STOP:
            break
        if label in SKIP_LABELS:
            continue
        for c in range(2, len(years) + 1):
            amount = ws.cell(row=r, column=c).value
            if not amount:
                continue
            year = years[c - 1]
            month = months[c - 1]
            if not year or not month:
                continue
            year_num = int(str(year).replace("年", ""))
            month_num = int(str(month).replace("月", ""))
            rows.append({
                "user_id": USER_ID,
                "year_month": f"{year_num:04d}-{month_num:02d}-01",
                "category": label,
                "kind": kind,
                "amount": float(amount),
            })
    insert_rows("monthly_entries", rows)


def import_investments(ws):
    blocks = [("NISA成長投資枠", 1, 2, 3), ("NISA積立投資枠", 6, 7, 8), ("スペースX", 11, 12, None)]
    rows = []
    for account, date_col, contrib_col, val_col in blocks:
        for r in range(3, ws.max_row + 1):
            date = ws.cell(row=r, column=date_col).value
            contrib = ws.cell(row=r, column=contrib_col).value
            if date is None or contrib is None:
                continue
            valuation = ws.cell(row=r, column=val_col).value if val_col else None
            rows.append({
                "user_id": USER_ID,
                "account": account,
                "entry_date": date.date().isoformat() if hasattr(date, "date") else str(date),
                "contribution": float(contrib),
                "valuation": float(valuation) if valuation else None,
            })
    insert_rows("investment_entries", rows)


def export_special_expenses_csv(ws, out_path):
    """特別出費シートは月ごとに列がずれるため、自動取り込みせずCSVに書き出して確認用にする"""
    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                writer.writerow(row)
    print(f"  -> 特別出費シートを {out_path} に書き出しました。内容を確認し、アプリのUIから手入力してください。")


def main():
    if len(sys.argv) != 2:
        print("使い方: python scripts/import_excel.py /path/to/家計簿.xlsx")
        sys.exit(1)
    if not (SUPABASE_URL and SERVICE_KEY and USER_ID):
        print("環境変数 SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY / SUPABASE_USER_ID を設定してください")
        sys.exit(1)

    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)

    print("生活管理 を取り込み中...")
    import_monthly_entries(wb["生活管理"])

    print("投資管理 を取り込み中...")
    import_investments(wb["投資管理"])

    print("特別出費 をCSVに書き出し中...")
    export_special_expenses_csv(wb["特別出費"], "special_expenses_export.csv")

    print("完了しました。")


if __name__ == "__main__":
    main()
